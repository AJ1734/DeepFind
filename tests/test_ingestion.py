import os
import pymupdf
import docx, openpyxl
from pptx import Presentation
from PIL import Image

from ingestion_service import IngestionService


def make_files(d):
    d.mkdir(exist_ok=True)
    (d / "a.txt").write_text("hello world zebra")
    x = docx.Document(); x.add_paragraph("word doc text"); x.save(d / "b.docx")
    wb = openpyxl.Workbook(); wb.active["A1"] = "excel cell"; wb.save(d / "c.xlsx")
    p = Presentation(); s = p.slides.add_slide(p.slide_layouts[5]); s.shapes.title.text = "slide title"; p.save(d / "d.pptx")
    pdf = pymupdf.open(); pdf.new_page().insert_text((72, 72), "digital pdf text is long enough"); pdf.save(d / "e.pdf")
    Image.new("RGB", (50, 50)).save(d / "f.png")
    # scanned PDF: one page that is just an image
    img = d / "scan_src.png"; Image.new("RGB", (200, 200), "white").save(img)
    pdf = pymupdf.open(); pg = pdf.new_page(); pg.insert_image(pg.rect, filename=str(img)); pdf.save(d / "scan.pdf")
    img.unlink()
    # mixed: page 1 digital, page 2 scanned image
    img2 = d / "s2.png"; Image.new("RGB", (200, 200), "white").save(img2)
    pdf = pymupdf.open(); pdf.new_page().insert_text((72, 72), "page one has real digital text here")
    pg = pdf.new_page(); pg.insert_image(pg.rect, filename=str(img2)); pdf.save(d / "mixed.pdf"); img2.unlink()
    (d / "node_modules").mkdir(); (d / "node_modules" / "junk.txt").write_text("ignore me")
    (d / "video.mp4").write_bytes(b"\x00" * 100)


def run(tmp_path, folder):
    svc = IngestionService(data_dir=tmp_path / "cache")
    got = []
    svc.register_appu_callback(got.append)
    svc.update_limits(100, 100, False)
    svc.start_watcher([str(folder)])
    assert svc.wait_idle(30)
    return svc, got


def test_backfill_extraction_and_routing(tmp_path):
    d = tmp_path / "w"; make_files(d)
    svc, got = run(tmp_path, d)
    try:
        by = {g["file_name"]: g for g in got}
        assert "zebra" in by["a.txt"]["content"] and "word doc text" in by["b.docx"]["content"]
        assert "excel cell" in by["c.xlsx"]["content"] and "slide title" in by["d.pptx"]["content"]
        assert "digital pdf" in by["e.pdf"]["content"]
        assert "video.mp4" in by and by["video.mp4"]["content"] == ""      # findable by name at once
        assert "junk.txt" not in by                                        # ignored folder
        q = {i["file_name"] for i in list(svc.get_tier2_queue().queue)}
        assert q == {"f.png", "scan.pdf", "mixed.pdf", "video.mp4"}      # videos: frames follow in tier 2
        assert "page one" in by["mixed.pdf"]["content"]                    # partial text searchable now
        assert "scan.pdf" not in by and "f.png" not in by
    finally:
        svc.stop_watcher()


def test_hash_only_committed_after_mark_indexed(tmp_path):
    d = tmp_path / "w"; d.mkdir(); (d / "a.txt").write_text("hello")
    svc, got = run(tmp_path, d); svc.stop_watcher()
    assert len(got) == 1 and len(svc.hash_cache) == 1        # tier 1: delivered => committed
    svc2, got2 = run(tmp_path, d); svc2.stop_watcher()
    assert got2 == []                                         # unchanged -> skipped

    # failed delivery must NOT commit (old bug: file was never retried)
    d2 = tmp_path / "w2"; d2.mkdir(); (d2 / "b.txt").write_text("hi")
    bad = IngestionService(data_dir=tmp_path / "cache2")
    bad.register_appu_callback(lambda doc: (_ for _ in ()).throw(RuntimeError("db down")))
    bad.update_limits(100, 100, False); bad.start_watcher([str(d2)]); bad.wait_idle(30); bad.stop_watcher()
    assert len(bad.hash_cache) == 0


def test_tier2_not_committed_until_marked(tmp_path):
    d = tmp_path / "w"; d.mkdir(); Image.new("RGB", (40, 40)).save(d / "i.png")
    svc, _ = run(tmp_path, d)
    item = svc.get_tier2_queue().get_nowait()
    assert len(svc.hash_cache) == 0
    svc.mark_indexed(item["file_path"]); svc.stop_watcher()
    assert len(svc.hash_cache) == 1


def test_live_events_delete_move_and_folder_delete(tmp_path):
    import shutil, time
    d = tmp_path / "w"; d.mkdir()
    svc, got = run(tmp_path, d)
    try:
        (d / "new.txt").write_text("live"); svc.wait_idle(10); time.sleep(1.5); svc.wait_idle(10)
        assert any(g.get("file_name") == "new.txt" for g in got)
        os.rename(d / "new.txt", d / "moved.txt"); time.sleep(1.5); svc.wait_idle(10)
        assert any(g.get("deleted") and g["file_path"].endswith("new.txt") for g in got)
        assert any(g.get("file_name") == "moved.txt" for g in got)
        (d / "sub").mkdir(); (d / "sub" / "x.txt").write_text("x"); time.sleep(1.5); svc.wait_idle(10)
        shutil.rmtree(d / "sub"); time.sleep(1.5); svc.wait_idle(10)
        assert any(g.get("deleted") and g.get("recursive") for g in got)
    finally:
        svc.stop_watcher()
