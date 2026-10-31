"""
ingestion_service.py
---------------------
HP's module: file watcher, system throttle monitor, text extraction, queue orchestration.

Public API (per team contract):
    start_watcher(paths)            stop_watcher()
    update_limits(cpu_limit, ram_limit, pause_on_battery)
    register_appu_callback(cb)      # Tier 1 docs (and delete events) pushed here
    get_tier2_queue()               # Tier 2 (images / scanned PDFs) for Aaling
    should_pause_tier2()
Extras added for integration:
    mark_indexed(file_path)         # call AFTER a doc is safely stored (commits the hash)
    get_stats(), reset_cache(), wait_idle()

Payload pushed to Appu / put on the Tier 2 queue:
    {file_path, file_name, file_type, content, vector(None), modified_time, needs_ocr}
Deletes:  {"file_path": str, "deleted": True}  (+ "recursive": True for a whole folder)
"""

from __future__ import annotations

import gc
import json
import logging
import os
import queue
import threading
import time
from pathlib import Path
from typing import Callable, Optional

import psutil
import xxhash
from watchdog.events import FileSystemEventHandler
from watchdog.observers import Observer

logger = logging.getLogger("ingestion_service")

# --------------------------------------------------------------------------
# Config
# --------------------------------------------------------------------------

TEMP_SUFFIXES = {".crdownload", ".part", ".tmp", ".download"}
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".py", ".rs", ".html",
                   ".js", ".ts", ".css", ".yaml", ".yml", ".toml", ".sh"}
DOC_EXTENSIONS = {".docx", ".xlsx", ".pptx"}
PDF_EXTENSION = ".pdf"
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tiff", ".tif", ".bmp"}
VIDEO_EXTENSIONS = {".mp4", ".mkv", ".mov", ".avi", ".webm", ".m4v"}   # not hashed: never read GBs
HASHED_EXTENSIONS = TEXT_EXTENSIONS | DOC_EXTENSIONS | {PDF_EXTENSION} | IMAGE_EXTENSIONS

# Never index dependency / build / cache folders (would drown the index in junk).
IGNORED_DIR_NAMES = {".git", "node_modules", ".venv", "venv", "env", "__pycache__",
                     ".idea", ".vscode", "target", "dist", "build", ".pytest_cache",
                     ".local_semantic_search", "$RECYCLE.BIN", "System Volume Information"}

MAX_TEXT_BYTES = 5 * 1024 * 1024        # plain-text read cap (8GB-RAM target)
MAX_CONTENT_CHARS = 1_000_000           # stored-content cap per file
TEXT_LAYER_MIN_CHARS = 20               # PDF page with fewer chars = "no text layer"
RECENT_FILE_SEC = 3.0                   # only debounce files modified this recently
DEBOUNCE_STABLE_CHECK_DELAY_SEC = 0.5
MONITOR_POLL_INTERVAL_SEC = 2.5
GC_EVERY_N_FILES = 25
HASH_SAVE_INTERVAL_SEC = 2.0

DATA_DIR = Path(os.environ.get("DEEPFIND_DATA_DIR", Path.home() / ".local_semantic_search"))


def _is_ignored(path: str) -> bool:
    p = Path(path)
    if p.name.startswith("~$") or p.suffix.lower() in TEMP_SUFFIXES:
        return True
    return any(part in IGNORED_DIR_NAMES for part in p.parts)


# --------------------------------------------------------------------------
# System health / throttle monitor
# --------------------------------------------------------------------------

class SystemMonitor:
    def __init__(self, cpu_limit: int = 50, ram_limit: int = 75, pause_on_battery: bool = True):
        self.cpu_limit, self.ram_limit, self.pause_on_battery = cpu_limit, ram_limit, pause_on_battery
        self.last_cpu = 0.0
        self.last_ram = 0.0
        self.on_battery = False
        self._pause_flag = threading.Event()
        self._stop_flag = threading.Event()
        self._thread: Optional[threading.Thread] = None
        self._lock = threading.Lock()

    def update_limits(self, cpu_limit: int, ram_limit: int, pause_on_battery: bool):
        with self._lock:
            self.cpu_limit, self.ram_limit, self.pause_on_battery = cpu_limit, ram_limit, pause_on_battery

    def start(self):
        if self._thread and self._thread.is_alive():
            return
        self._stop_flag.clear()
        self._thread = threading.Thread(target=self._run, daemon=True, name="system-monitor")
        self._thread.start()

    def stop(self):
        self._stop_flag.set()
        if self._thread:
            self._thread.join(timeout=5)

    def should_pause_tier2(self) -> bool:
        return self._pause_flag.is_set()

    def _run(self):
        psutil.cpu_percent(interval=None)  # warm-up call
        while not self._stop_flag.is_set():
            with self._lock:
                cpu_limit, ram_limit, on_batt_rule = self.cpu_limit, self.ram_limit, self.pause_on_battery
            cpu_pct = psutil.cpu_percent(interval=None)
            ram_pct = psutil.virtual_memory().percent
            on_battery = False
            try:
                batt = psutil.sensors_battery()
                if batt is not None:
                    on_battery = not batt.power_plugged
            except Exception:
                pass
            self.last_cpu, self.last_ram, self.on_battery = cpu_pct, ram_pct, on_battery
            if cpu_pct > cpu_limit or ram_pct > ram_limit or (on_batt_rule and on_battery):
                self._pause_flag.set()
            else:
                self._pause_flag.clear()
            self._stop_flag.wait(MONITOR_POLL_INTERVAL_SEC)


# --------------------------------------------------------------------------
# Hash cache — a hash is only COMMITTED after the doc was actually indexed
# --------------------------------------------------------------------------

class HashCache:
    def __init__(self, path: Optional[Path] = None):
        self.path = Path(path) if path else DATA_DIR / "hash_cache.json"
        self._cache: dict[str, str] = {}
        self._pending: dict[str, str] = {}   # computed, not yet committed
        self._stat: dict[str, tuple] = {}    # path -> (size, mtime_ns) when indexed; this run only
        self._lock = threading.Lock()
        self._dirty = False
        self._last_save = 0.0
        self._load()

    def _load(self):
        try:
            if self.path.exists():
                self._cache = json.loads(self.path.read_text())
        except Exception as e:
            logger.warning("Could not load hash cache: %s", e)
            self._cache = {}

    def _save_locked(self, force: bool = False):
        if not self._dirty:
            return
        if not force and time.time() - self._last_save < HASH_SAVE_INTERVAL_SEC:
            return
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            tmp = self.path.with_suffix(".tmp")
            tmp.write_text(json.dumps(self._cache))
            os.replace(tmp, self.path)
            self._dirty = False
            self._last_save = time.time()
        except Exception as e:
            logger.warning("Could not save hash cache: %s", e)

    @staticmethod
    def hash_file(file_path: str) -> str:
        h = xxhash.xxh64()
        with open(file_path, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                h.update(chunk)
        return h.hexdigest()

    @classmethod
    def signature(cls, file_path: str, ext: str) -> str:
        if ext in HASHED_EXTENSIONS:
            return cls.hash_file(file_path)
        st = os.stat(file_path)  # unsupported types (videos etc.): don't read GBs
        return f"{st.st_size}:{st.st_mtime_ns}"

    def is_unchanged(self, file_path: str, ext: str = "") -> bool:
        # Fast path for repeated scans (scheduled runs): same size + modified time as when this
        # run indexed it -> unchanged, without reading the file again.
        try:
            st = os.stat(file_path)
            with self._lock:
                if file_path in self._cache and self._stat.get(file_path) == (st.st_size, st.st_mtime_ns):
                    return True
        except OSError:
            return False
        try:
            sig = self.signature(file_path, ext or Path(file_path).suffix.lower())
        except OSError:
            return False
        with self._lock:
            self._pending[file_path] = sig
            return self._cache.get(file_path) == sig

    def mark_indexed(self, file_path: str):
        with self._lock:
            sig = self._pending.pop(file_path, None)
        if sig is None:
            try:
                sig = self.signature(file_path, Path(file_path).suffix.lower())
            except OSError:
                return
        try:
            st = os.stat(file_path)
            stat = (st.st_size, st.st_mtime_ns)
        except OSError:
            stat = None
        with self._lock:
            self._cache[file_path] = sig
            if stat:
                self._stat[file_path] = stat
            self._dirty = True
            self._save_locked()

    def purge(self, file_path: str):
        with self._lock:
            self._pending.pop(file_path, None)
            self._stat.pop(file_path, None)
            if self._cache.pop(file_path, None) is not None:
                self._dirty = True
                self._save_locked()

    def purge_prefix(self, folder: str):
        prefix = folder.rstrip("/\\") + os.sep
        with self._lock:
            for k in [k for k in self._cache if k.startswith(prefix)]:
                del self._cache[k]
                self._stat.pop(k, None)
                self._dirty = True
            self._save_locked()

    def purge_extensions(self, exts) -> int:
        """Forget the hashes of every file with one of these extensions, so the next scan
        re-reads them (used when the OCR engine changes). Returns how many were forgotten."""
        exts = {e.lower() for e in exts}
        with self._lock:
            keys = [k for k in self._cache if Path(k).suffix.lower() in exts]
            for k in keys:
                del self._cache[k]
                self._stat.pop(k, None)
            if keys:
                self._dirty = True
                self._save_locked(force=True)
        return len(keys)

    def clear(self):
        with self._lock:
            self._cache.clear()
            self._pending.clear()
            self._stat.clear()
            self._dirty = True
            self._save_locked(force=True)

    def flush(self):
        with self._lock:
            self._save_locked(force=True)

    def __len__(self):
        return len(self._cache)


# --------------------------------------------------------------------------
# Text extraction (Tier 1)
# --------------------------------------------------------------------------

def _extract_pdf(file_path: str) -> tuple[str, bool]:
    """(text, needs_ocr). needs_ocr = any page has no text layer AND looks scanned
    (contains an image / drawings). Truly blank pages don't trigger OCR."""
    import pymupdf as fitz

    parts, needs_ocr = [], False
    with fitz.open(file_path) as doc:
        for page in doc:
            if len(parts) >= 100:  # Prevent OOM on massive PDFs
                break
            txt = page.get_text().strip()
            if len(txt) >= TEXT_LAYER_MIN_CHARS:
                parts.append(txt)
                continue
            if txt:
                parts.append(txt)
            if page.get_images() or page.get_drawings():
                needs_ocr = True
    return "\n".join(parts), needs_ocr


def _extract_docx(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_xlsx(file_path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    try:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        parts.append(str(cell))
    finally:
        wb.close()
    return "\n".join(parts)


def _extract_pptx(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    parts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text:
                            parts.append(run.text)
    return "\n".join(parts)


def _extract_plain_text(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read(MAX_TEXT_BYTES)


def extract_text(file_path: str, ext: str) -> tuple[str, bool]:
    """Returns (content, needs_ocr)."""
    try:
        if ext == PDF_EXTENSION:
            text, needs_ocr = _extract_pdf(file_path)
        elif ext == ".docx":
            text, needs_ocr = _extract_docx(file_path), False
        elif ext == ".xlsx":
            text, needs_ocr = _extract_xlsx(file_path), False
        elif ext == ".pptx":
            text, needs_ocr = _extract_pptx(file_path), False
        elif ext in TEXT_EXTENSIONS:
            text, needs_ocr = _extract_plain_text(file_path), False
        elif ext in IMAGE_EXTENSIONS or ext in VIDEO_EXTENSIONS:
            return "", True
        else:
            return "", False   # unsupported: still indexed by file name
        return text[:MAX_CONTENT_CHARS], needs_ocr
    except Exception as e:
        logger.warning("Failed to extract %s: %s", file_path, e)
        return "", False


def _wait_until_stable(file_path: str, checks: int = 10,
                       delay: float = DEBOUNCE_STABLE_CHECK_DELAY_SEC) -> bool:
    """Wait until size stops changing (file fully written). Old files skip the wait."""
    try:
        if time.time() - os.path.getmtime(file_path) > RECENT_FILE_SEC:
            return True
        last = -1
        for _ in range(checks):
            size = os.path.getsize(file_path)
            if size == last:
                return True
            last = size
            time.sleep(delay)
        return True
    except OSError:
        return False


# --------------------------------------------------------------------------
# Core service
# --------------------------------------------------------------------------

class IngestionService:
    def __init__(self, data_dir: Optional[Path] = None):
        self.monitor = SystemMonitor()
        self.hash_cache = HashCache(Path(data_dir) / "hash_cache.json" if data_dir else None)
        self._observer: Optional[Observer] = None
        self._appu_callback: Optional[Callable[[dict], None]] = None
        self._tier2_queue: "queue.Queue[dict]" = queue.Queue()
        self._watched_paths: list[str] = []
        # single worker thread: watchdog callbacks only enqueue, never block/process
        self._work_q: "queue.Queue[Optional[tuple]]" = queue.Queue()
        self._pending: set[str] = set()
        self._pending_lock = threading.Lock()
        self._worker: Optional[threading.Thread] = None
        self._backfill_done = threading.Event()
        self._backfill_done.set()
        self._current: Optional[str] = None
        self._processed = 0
        self.errors = 0
        # Scheduled indexing: when _defer is on, file events are only remembered (cheap) and
        # read later by index_now(). Deletions are always applied straight away.
        self._defer = False
        self._deferred: set[str] = set()
        self._burst = 0                      # >0 while index_now() is feeding the queue
        self._burst_lock = threading.Lock()

    # ---- public API ------------------------------------------------------

    def start_watcher(self, paths: list[str], scan: bool = True):
        """Watch `paths`. scan=False skips the start-up walk (scheduled / manual indexing)."""
        self.stop_watcher()
        self._drain_tier2()   # stale items from old folders must never be processed
        self._watched_paths = [os.path.normpath(os.path.abspath(p)) for p in paths]
        with self._pending_lock:   # forget remembered changes from folders no longer watched
            self._deferred = {p for p in self._deferred if self.is_watched(p)}
        self.monitor.start()
        self._start_worker()

        handler = _Handler(self)
        self._observer = Observer()
        valid = []
        for p in self._watched_paths:
            if os.path.isdir(p):
                self._observer.schedule(handler, p, recursive=True)
                valid.append(p)
            else:
                logger.warning("Path does not exist, skipping: %s", p)
        self._observer.start()
        logger.info("Watching: %s", valid)

        if scan:
            self._backfill_done.clear()
            threading.Thread(target=self._backfill_scan, args=(valid,), daemon=True,
                             name="backfill").start()

    def stop_watcher(self):
        if self._observer:
            self._observer.stop()
            self._observer.join(timeout=5)
            self._observer = None
        self._stop_worker()
        self.monitor.stop()
        self.hash_cache.flush()

    def _drain_tier2(self):
        try:
            while True:
                self._tier2_queue.get_nowait()
        except queue.Empty:
            pass

    # ---- scheduled indexing ---------------------------------------------------

    def set_deferred(self, on: bool):
        """on=True: remember changes instead of reading them. on=False: back to instant,
        and anything remembered meanwhile is read now."""
        was = self._defer
        self._defer = bool(on)
        if was and not on:
            self.index_now()

    def pending_changes(self) -> int:
        return len(self._deferred)

    def index_now(self, paths: Optional[list[str]] = None, exts: Optional[set] = None):
        """Read what changed: everything remembered by the watcher PLUS a walk over the watched
        folders (or just `paths`), which also catches files that appeared or changed while the app
        was closed. Unchanged files are skipped by the hash cache. `exts` limits to some types."""
        if not (self._worker and self._worker.is_alive()):
            return      # not watching right now: start_watcher() does its own full scan
        with self._burst_lock:
            self._burst += 1
        with self._pending_lock:
            remembered, self._deferred = list(self._deferred), set()
        roots = [os.path.normpath(os.path.abspath(p)) for p in (paths if paths is not None else self._watched_paths)]
        self._backfill_done.clear()
        threading.Thread(target=self._scan_now, args=(remembered, exts, roots), daemon=True,
                         name="index-now").start()

    def _scan_now(self, remembered: list[str], exts: Optional[set], roots: Optional[list[str]] = None):
        try:
            for p in remembered:
                if exts is None or Path(p).suffix.lower() in exts:
                    self.submit_file(p)
            for root in roots or []:
                if os.path.isdir(root):
                    self.submit_dir(root, exts)
                elif os.path.isfile(root) and (exts is None or Path(root).suffix.lower() in exts):
                    self.submit_file(root)
        except Exception:
            self.errors += 1
            logger.exception("index_now scan failed")
        finally:
            with self._burst_lock:
                self._burst -= 1
            self._backfill_done.set()

    def is_watched(self, file_path: str) -> bool:
        fp = os.path.normpath(file_path)
        return any(fp.startswith(root.rstrip("/\\") + os.sep) for root in self._watched_paths)

    def update_limits(self, cpu_limit: int, ram_limit: int, pause_on_battery: bool):
        self.monitor.update_limits(cpu_limit, ram_limit, pause_on_battery)

    def register_appu_callback(self, callback: Callable[[dict], None]):
        self._appu_callback = callback

    def get_tier2_queue(self) -> "queue.Queue[dict]":
        return self._tier2_queue

    def should_pause_tier2(self) -> bool:
        return self.monitor.should_pause_tier2()

    def mark_indexed(self, file_path: str):
        self.hash_cache.mark_indexed(file_path)

    def reset_cache(self):
        self.hash_cache.clear()

    def get_stats(self) -> dict:
        m = self.monitor
        return {
            "cpu_percent": m.last_cpu, "ram_percent": m.last_ram, "on_battery": m.on_battery,
            "paused": m.should_pause_tier2(), "indexed_files": len(self.hash_cache),
            "tier2_queue": self._tier2_queue.qsize(), "work_queue": self._work_q.qsize(),
            "current_file": self._current, "watching": list(self._watched_paths),
            "errors": self.errors, "pending_changes": len(self._deferred),
        }

    def wait_idle(self, timeout: float = 30.0) -> bool:
        """Block until backfill finished and the work queue is drained (for tests/tools)."""
        end = time.time() + timeout
        while time.time() < end:
            if self._backfill_done.is_set() and self._work_q.unfinished_tasks == 0:
                return True
            time.sleep(0.05)
        return False

    # ---- event entry points (called from watchdog thread — enqueue only) -----

    def submit_file(self, file_path: str):
        if _is_ignored(file_path):
            return
        if self._defer and self._burst == 0:
            with self._pending_lock:       # scheduled mode: just remember it for the next run
                self._deferred.add(file_path)
            return
        with self._pending_lock:
            if file_path in self._pending:
                return   # already queued: collapses bursts of modified events
            self._pending.add(file_path)
        self._work_q.put(("upsert", file_path))

    def submit_dir(self, dir_path: str, exts: Optional[set] = None):
        for dirpath, dirnames, filenames in os.walk(dir_path):
            dirnames[:] = [d for d in dirnames if d not in IGNORED_DIR_NAMES]
            for fname in filenames:
                if exts is None or Path(fname).suffix.lower() in exts:
                    self.submit_file(os.path.join(dirpath, fname))

    def submit_delete(self, path: str, recursive: bool = False):
        with self._pending_lock:           # a remembered change to a file that is now gone
            if recursive:
                prefix = path.rstrip("/\\") + os.sep
                self._deferred = {p for p in self._deferred if not p.startswith(prefix)}
            else:
                self._deferred.discard(path)
        self._work_q.put(("delete_dir" if recursive else "delete", path))

    # ---- worker -------------------------------------------------------------

    def _start_worker(self):
        if self._worker and self._worker.is_alive():
            return
        self._worker = threading.Thread(target=self._run_worker, daemon=True, name="ingest-worker")
        self._worker.start()

    def _stop_worker(self):
        if not (self._worker and self._worker.is_alive()):
            return
        try:   # drop queued work, then send the stop sentinel
            while True:
                self._work_q.get_nowait()
                self._work_q.task_done()
        except queue.Empty:
            pass
        with self._pending_lock:
            self._pending.clear()
        self._work_q.put(None)
        self._worker.join(timeout=10)
        self._worker = None
        self._backfill_done.set()

    def _run_worker(self):
        while True:
            item = self._work_q.get()
            try:
                if item is None:
                    return
                kind, path = item
                if kind == "delete":
                    self._handle_deleted(path)
                elif kind == "delete_dir":
                    self._handle_deleted(path, recursive=True)
                else:
                    with self._pending_lock:
                        self._pending.discard(path)   # events during processing re-queue
                    self._process_file(path)
            except Exception:
                self.errors += 1
                logger.exception("ingestion worker error on %r", item)
            finally:
                self._current = None
                self._work_q.task_done()

    def _backfill_scan(self, paths: list[str]):
        try:
            for root in paths:
                self.submit_dir(root)
        finally:
            self._backfill_done.set()

    def _deliver(self, doc: dict) -> bool:
        cb = self._appu_callback
        if cb is None:
            return False
        try:
            cb(doc)
            return True
        except Exception:
            self.errors += 1
            logger.exception("Appu callback failed for %s", doc.get("file_path"))
            return False

    def _handle_deleted(self, path: str, recursive: bool = False):
        if recursive:
            self.hash_cache.purge_prefix(path)
            self._deliver({"file_path": path, "deleted": True, "recursive": True})
        else:
            self.hash_cache.purge(path)
            self._deliver({"file_path": path, "deleted": True})

    def _process_file(self, file_path: str):
        ext = Path(file_path).suffix.lower()
        if _is_ignored(file_path) or not os.path.isfile(file_path):
            return
        self._current = file_path
        if not _wait_until_stable(file_path):
            return
        if self.hash_cache.is_unchanged(file_path, ext):
            return

        content, needs_ocr = extract_text(file_path, ext)
        doc = {
            "file_path": file_path,
            "file_name": Path(file_path).name,
            "file_type": ext.lstrip("."),
            "content": content,
            "vector": None,
            "modified_time": os.path.getmtime(file_path),
            "needs_ocr": needs_ocr,
        }

        if needs_ocr:
            if content or ext in VIDEO_EXTENSIONS:   # searchable by name now, text/frames follow later
                self._deliver(dict(doc, needs_ocr=False))
            self._tier2_queue.put(doc)     # consumer gates on should_pause_tier2()
        elif self._deliver(doc):
            self.hash_cache.mark_indexed(file_path)

        self._processed += 1
        if self._processed % GC_EVERY_N_FILES == 0:
            gc.collect()


class _Handler(FileSystemEventHandler):
    def __init__(self, service: IngestionService):
        self.s = service

    def on_created(self, event):
        (self.s.submit_dir if event.is_directory else self.s.submit_file)(event.src_path)

    def on_modified(self, event):
        if not event.is_directory:
            self.s.submit_file(event.src_path)

    def on_deleted(self, event):
        self.s.submit_delete(event.src_path, recursive=event.is_directory)

    def on_moved(self, event):
        self.s.submit_delete(event.src_path, recursive=event.is_directory)
        (self.s.submit_dir if event.is_directory else self.s.submit_file)(event.dest_path)


# --------------------------------------------------------------------------
# Module-level singleton + public API
# --------------------------------------------------------------------------

_service = IngestionService()


def start_watcher(paths: list[str], scan: bool = True): _service.start_watcher(paths, scan)
def stop_watcher(): _service.stop_watcher()
def update_limits(cpu_limit: int, ram_limit: int, pause_on_battery: bool):
    _service.update_limits(cpu_limit, ram_limit, pause_on_battery)
def register_appu_callback(callback: Callable[[dict], None]): _service.register_appu_callback(callback)
def get_tier2_queue() -> "queue.Queue[dict]": return _service.get_tier2_queue()
def should_pause_tier2() -> bool: return _service.should_pause_tier2()
def mark_indexed(file_path: str): _service.mark_indexed(file_path)
def reset_cache(): _service.reset_cache()
def get_stats() -> dict: return _service.get_stats()
def wait_idle(timeout: float = 30.0) -> bool: return _service.wait_idle(timeout)
def purge_hash_prefix(folder: str): _service.hash_cache.purge_prefix(folder)
def is_watched(file_path: str) -> bool: return _service.is_watched(file_path)
def set_deferred(on: bool): _service.set_deferred(on)
def index_now(paths: Optional[list[str]] = None, exts: Optional[set] = None): _service.index_now(paths, exts)
def purge_hash_extensions(exts) -> int: return _service.hash_cache.purge_extensions(exts)


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)

    def fake_appu(doc: dict):
        print(f"[Appu] {doc.get('file_name', doc.get('file_path'))} "
              f"({len(doc.get('content', ''))} chars){' [DELETED]' if doc.get('deleted') else ''}")

    register_appu_callback(fake_appu)
    update_limits(cpu_limit=50, ram_limit=75, pause_on_battery=True)
    os.makedirs("./test_watch_folder", exist_ok=True)
    start_watcher(["./test_watch_folder"])
    print("Watching ./test_watch_folder — Ctrl+C to stop.")
    try:
        while True:
            time.sleep(1)
            while not get_tier2_queue().empty():
                print(f"[Tier2] {get_tier2_queue().get()['file_name']} needs OCR/vision")
    except KeyboardInterrupt:
        stop_watcher()